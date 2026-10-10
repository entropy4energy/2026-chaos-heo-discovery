#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PCA analysis of the CHAOS descriptor space colored/evaluated by EFA.

This script uses two descriptor blocks as PCA inputs:
  - H-L: DFT-derived lattice-distortion descriptors
  - S-AK: non-DFT chemistry descriptors

It intentionally excludes EFA and related target columns
from the PCA input matrix. EFA is used only after PCA to compute percentile ranks
and PC-EFA correlations for coloring and interpretation.

Outputs
-------
1. Excel workbook containing PCA scores, loadings, explained variance, and plot data.
2. TXT report explaining the calculation process and key results.

Local dependencies
------------------
pip install pandas numpy scipy scikit-learn openpyxl
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

import numpy as np
import pandas as pd
from scipy import stats
from sklearn.decomposition import PCA
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import StandardScaler


DFT_PRETTY_LABELS: Dict[str, str] = {
    "lat_distortion_misfit_mean": r"$\bar{\epsilon}$",
    "lat_distortion_misfit_stdev": r"$\sigma_{\epsilon}$",
    "lat_distortion_relax_distance_mean": r"$\bar{d}$",
    "lat_distortion_relax_distance_stdev": r"$\sigma_d$",
    "solubility_parameter": r"$\delta_s$",
}

CHEM_PRETTY_LABELS: Dict[str, str] = {
    "VEC": "VEC",
    "delta_electronegativity_Allen": r"$\delta\chi_{\mathrm{Allen}}$",
    "delta_electronegativity_Ghosh": r"$\delta\chi_{\mathrm{Ghosh}}$",
    "delta_electronegativity_Pauling": r"$\delta\chi_{\mathrm{Pauling}}$",
    "delta_electronegativity_Pearson": r"$\delta\chi_{\mathrm{Pearson}}$",
    "delta_radii_Ghosh08": r"$\delta r_{\mathrm{Ghosh}}$",
    "delta_radii_Pyykko": r"$\delta r_{\mathrm{Pyykk\ddot{o}}}$",
    "delta_radii_Slatter": r"$\delta r_{\mathrm{Slater}}$",
    "delta_radius_PT": r"$\delta r_{\mathrm{PT}}$",
    "delta_radius_Saxena": r"$\delta r_{\mathrm{Saxena}}$",
    "delta_radius_covalent_PT": r"$\delta r_{\mathrm{cov,PT}}$",
    "delta_radius_covalent": r"$\delta r_{\mathrm{cov}}$",
    "param_geo_radii_Ghosh08": r"$\Lambda_{\mathrm{Ghosh}}$",
    "param_geo_radii_Pyykko": r"$\Lambda_{\mathrm{Pyykk\ddot{o}}}$",
    "param_geo_radii_Slatter": r"$\Lambda_{\mathrm{Slater}}$",
    "param_geo_radius_PT": r"$\Lambda_{\mathrm{PT}}$",
    "param_geo_radius_Saxena": r"$\Lambda_{\mathrm{Saxena}}$",
    "param_geo_radius_covalent_PT": r"$\Lambda_{\mathrm{cov,PT}}$",
    "param_geo_radius_covalent": r"$\Lambda_{\mathrm{cov}}$",
}

PRETTY_LABELS = {**DFT_PRETTY_LABELS, **CHEM_PRETTY_LABELS}

EXCLUDED_IF_PRESENT = {
    "EFA", "efa_original_row", "match_count",
    "entropy_forming_ability", "enthalpy_mix_atom",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="PCA on H-L DFT descriptors + S-AK chemistry descriptors, colored/evaluated by EFA."
    )
    parser.add_argument("--input", required=True, help="Input Excel workbook, e.g. locked_493.xlsx")
    parser.add_argument("--sheet", default="combined", help="Input sheet name. Default: combined")
    parser.add_argument("--outdir", default="pca_efa_chaos_descriptor_space_outputs", help="Output directory")
    parser.add_argument("--efa-col", default="EFA", help="EFA column name. Default: EFA")
    parser.add_argument("--dft-start", default="H", help="Start Excel column for DFT descriptors. Default: H")
    parser.add_argument("--dft-end", default="L", help="End Excel column for DFT descriptors. Default: L")
    parser.add_argument("--chem-start", default="S", help="Start Excel column for chemistry descriptors. Default: S")
    parser.add_argument("--chem-end", default="AK", help="End Excel column for chemistry descriptors. Default: AK")
    parser.add_argument("--max-components", type=int, default=10, help="Maximum PCA components to output. Default: 10")
    parser.add_argument(
        "--family-cols",
        nargs="*",
        default=["Metal1", "Metal2", "Metal3", "Metal4"],
        help="Columns used to define compositional family if present.",
    )
    parser.add_argument(
        "--orient-pc2-positive-efa",
        action="store_true",
        default=True,
        help="Orient PC2 so that Spearman(PC2, EFA percentile) is positive. PCA signs are arbitrary; default is on.",
    )
    return parser.parse_args()


def excel_col_to_index0(col: str) -> int:
    col = col.strip().upper()
    idx = 0
    for ch in col:
        if not ("A" <= ch <= "Z"):
            raise ValueError(f"Invalid Excel column: {col}")
        idx = idx * 26 + (ord(ch) - ord("A") + 1)
    return idx - 1


def percentile_rank(series: pd.Series) -> pd.Series:
    s = pd.to_numeric(series, errors="coerce")
    out = pd.Series(np.nan, index=s.index, dtype=float)
    valid = s.dropna()
    n = len(valid)
    if n == 0:
        return out
    if n == 1:
        out.loc[valid.index] = 0.5
        return out
    ranks = valid.rank(method="average", ascending=True)
    out.loc[valid.index] = (ranks - 1) / (n - 1)
    return out


def quartile_from_percentile(p: pd.Series) -> pd.Series:
    labels = []
    for val in p:
        if pd.isna(val):
            labels.append(np.nan)
        elif val < 0.25:
            labels.append("Q1")
        elif val < 0.50:
            labels.append("Q2")
        elif val < 0.75:
            labels.append("Q3")
        else:
            labels.append("Q4")
    return pd.Series(labels, index=p.index)


def make_family(row: pd.Series, family_cols: Iterable[str]) -> str:
    vals = []
    for c in family_cols:
        if c in row.index and pd.notna(row[c]):
            s = str(row[c]).strip()
            if s and s.lower() != "nan":
                vals.append(s)
    return "-".join(sorted(vals)) if vals else "unknown_family"


def corr_pair(x: pd.Series, y: pd.Series, x_name: str, y_name: str) -> pd.DataFrame:
    temp = pd.DataFrame({x_name: pd.to_numeric(x, errors="coerce"), y_name: pd.to_numeric(y, errors="coerce")}).dropna()
    if len(temp) < 3:
        rows = [
            {"x": x_name, "y": y_name, "method": "Pearson", "N": len(temp), "coefficient": np.nan, "p_value": np.nan},
            {"x": x_name, "y": y_name, "method": "Spearman", "N": len(temp), "coefficient": np.nan, "p_value": np.nan},
        ]
    else:
        pr = stats.pearsonr(temp[x_name], temp[y_name])
        sr = stats.spearmanr(temp[x_name], temp[y_name])
        rows = [
            {"x": x_name, "y": y_name, "method": "Pearson", "N": len(temp), "coefficient": float(pr.statistic), "p_value": float(pr.pvalue)},
            {"x": x_name, "y": y_name, "method": "Spearman", "N": len(temp), "coefficient": float(sr.statistic), "p_value": float(sr.pvalue)},
        ]
    return pd.DataFrame(rows)


def autosize_workbook(path: Path) -> None:
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        wb = load_workbook(path)
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="D9E2F3")
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = Border(bottom=thin)
            for col_cells in ws.columns:
                col_letter = col_cells[0].column_letter
                max_len = 0
                for cell in col_cells[:1000]:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 42)
        wb.save(path)
    except Exception as exc:  # formatting is optional
        print(f"Warning: workbook formatting skipped: {exc}")


def write_report(
    out_txt: Path,
    input_path: Path,
    sheet: str,
    n_raw: int,
    n_used: int,
    dft_features: List[str],
    chem_features: List[str],
    valid_features: List[str],
    dropped_features: List[str],
    explained_df: pd.DataFrame,
    pc_corr_df: pd.DataFrame,
    loading_df: pd.DataFrame,
    out_xlsx: Path,
) -> None:
    def pc_corr_line(pc: str) -> str:
        rows = pc_corr_df[(pc_corr_df["PC"] == pc) & (pc_corr_df["target"] == "EFA_percentile")]
        pearson = rows[rows["method"] == "Pearson"]["coefficient"].iloc[0]
        spearman = rows[rows["method"] == "Spearman"]["coefficient"].iloc[0]
        return f"   {pc} vs EFA percentile: Pearson r = {pearson:.3f}; Spearman rho = {spearman:.3f}"

    lines = []
    lines.append("PCA of combined CHAOS descriptor space colored/evaluated by EFA")
    lines.append("=" * 78)
    lines.append("")
    lines.append("1. Data source")
    lines.append(f"   Input workbook: {input_path}")
    lines.append(f"   Sheet: {sheet}")
    lines.append(f"   Raw rows: {n_raw}")
    lines.append(f"   Rows used after dropping missing EFA: {n_used}")
    lines.append("")
    lines.append("2. PCA input features")
    lines.append("   DFT descriptor block: H-L")
    for f in dft_features:
        lines.append(f"      - {f} -> {PRETTY_LABELS.get(f, f)}")
    lines.append("   Non-DFT chemistry descriptor block: S-AK")
    for f in chem_features:
        lines.append(f"      - {f} -> {PRETTY_LABELS.get(f, f)}")
    lines.append(f"   Valid features used: {len(valid_features)}")
    if dropped_features:
        lines.append(f"   Dropped features: {', '.join(dropped_features)}")
    else:
        lines.append("   Dropped features: none")
    lines.append("")
    lines.append("3. Excluded variables")
    lines.append("   EFA was not used as a PCA input; it was used only for percentile coloring and post-PCA correlation.")
    lines.append("   Related target columns were not used.")
    lines.append("   Configurational entropy columns M-R were not used because the requested feature blocks are H-L and S-AK.")
    lines.append("")
    lines.append("4. Preprocessing")
    lines.append("   All selected features were converted to numeric values. Missing feature values were median-imputed.")
    lines.append("   Features were standardized to zero mean and unit variance before PCA.")
    lines.append("   PC2 was oriented to have a positive Spearman correlation with EFA percentile; PCA signs are arbitrary.")
    lines.append("")
    lines.append("5. Explained variance")
    for _, row in explained_df.head(5).iterrows():
        lines.append(
            f"   {row['component']}: explained variance ratio = {row['explained_variance_ratio']:.4f}; "
            f"cumulative = {row['cumulative_explained_variance_ratio']:.4f}"
        )
    lines.append("")
    lines.append("6. PC-EFA correlations")
    lines.append(pc_corr_line("PC1"))
    lines.append(pc_corr_line("PC2"))
    lines.append("")
    lines.append("7. Top absolute loadings")
    for pc in ["PC1", "PC2"]:
        col = f"abs_{pc}_loading"
        lines.append(f"   {pc}:")
        top = loading_df.sort_values(col, ascending=False).head(10)
        for _, r in top.iterrows():
            lines.append(
                f"      {r['feature_label']} ({r['feature_column']}; {r['feature_group']}): "
                f"loading = {r[f'{pc}_loading']:.4f}"
            )
    lines.append("")
    lines.append("8. Interpretation")
    lines.append("   The PC1-PC2 projection combines DFT lattice distortion and chemistry descriptors in one latent CHAOS descriptor space.")
    lines.append("   A systematic EFA gradient in this plane indicates that high-EFA systems are organized by the descriptor space, not by EFA target leakage.")
    lines.append("   The loading plots identify whether the EFA-sensitive latent direction is dominated by structural distortion metrics, chemistry-only metrics, or both.")
    lines.append("")
    lines.append("9. Output")
    lines.append(f"   Excel workbook: {out_xlsx}")
    lines.append(f"   TXT report: {out_txt}")
    out_txt.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    args = parse_args()
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    input_path = Path(args.input)
    if not input_path.exists():
        raise FileNotFoundError(f"Input workbook not found: {input_path}")

    df = pd.read_excel(input_path, sheet_name=args.sheet)
    n_raw = len(df)
    if args.efa_col not in df.columns:
        raise ValueError(f"EFA column not found: {args.efa_col}")

    dft_start = excel_col_to_index0(args.dft_start)
    dft_end = excel_col_to_index0(args.dft_end)
    chem_start = excel_col_to_index0(args.chem_start)
    chem_end = excel_col_to_index0(args.chem_end)
    n_cols = len(df.columns)
    for label, start, end in [("DFT", dft_start, dft_end), ("chemistry", chem_start, chem_end)]:
        if start < 0 or end >= n_cols or start > end:
            raise ValueError(f"Invalid {label} range for workbook with {n_cols} columns")

    dft_features = list(df.columns[dft_start:dft_end + 1])
    chem_features = list(df.columns[chem_start:chem_end + 1])
    candidate_features = dft_features + chem_features

    forbidden = [c for c in candidate_features if c in EXCLUDED_IF_PRESENT]
    if forbidden:
        raise ValueError("Selected PCA feature blocks contain excluded target columns: " + ", ".join(forbidden))

    data = df.copy()
    data[args.efa_col] = pd.to_numeric(data[args.efa_col], errors="coerce")
    data = data.dropna(subset=[args.efa_col]).reset_index(drop=True)
    n_used = len(data)
    if n_used < 3:
        raise ValueError("At least three rows with EFA values are required.")

    X_raw = data[candidate_features].apply(pd.to_numeric, errors="coerce")

    valid_features = []
    dropped_features = []
    for c in candidate_features:
        s = X_raw[c]
        if s.notna().sum() == 0 or s.dropna().nunique() <= 1:
            dropped_features.append(c)
        else:
            valid_features.append(c)
    if len(valid_features) < 2:
        raise ValueError("At least two valid PCA features are required.")

    feature_group = {c: "DFT" for c in dft_features}
    feature_group.update({c: "chemistry" for c in chem_features})

    X_valid = X_raw[valid_features]
    imputer = SimpleImputer(strategy="median")
    scaler = StandardScaler()
    X_imputed = imputer.fit_transform(X_valid)
    X_scaled = scaler.fit_transform(X_imputed)

    n_components = min(args.max_components, X_scaled.shape[0], X_scaled.shape[1])
    pca = PCA(n_components=n_components, random_state=0)
    scores = pca.fit_transform(X_scaled)
    components = pca.components_.copy()

    efa_percentile = percentile_rank(data[args.efa_col])

    # Orient PC2 to positive EFA correlation for easier interpretation. PCA signs are arbitrary.
    orientation_records = []
    for pc_idx in range(min(2, n_components)):
        pc_name = f"PC{pc_idx+1}"
        rho = stats.spearmanr(scores[:, pc_idx], efa_percentile, nan_policy="omit").statistic
        sign_multiplier = 1.0
        if pc_name == "PC2" and args.orient_pc2_positive_efa and np.isfinite(rho) and rho < 0:
            sign_multiplier = -1.0
            scores[:, pc_idx] *= -1.0
            components[pc_idx, :] *= -1.0
            rho = -rho
        orientation_records.append({"component": pc_name, "sign_multiplier_applied": sign_multiplier, "spearman_with_EFA_percentile_after_orientation": float(rho) if np.isfinite(rho) else np.nan})

    data["EFA_percentile"] = efa_percentile
    data["EFA_quartile"] = quartile_from_percentile(efa_percentile)
    data["anonymous_id"] = [f"CHAOS_{i+1:04d}" for i in range(n_used)]
    data["compositional_family"] = data.apply(lambda r: make_family(r, args.family_cols), axis=1)

    pc_cols = [f"PC{i+1}" for i in range(n_components)]
    score_df = pd.DataFrame(scores, columns=pc_cols)

    scores_public = pd.concat([
        data[["anonymous_id", "compositional_family", "EFA_percentile", "EFA_quartile"]].reset_index(drop=True),
        score_df.reset_index(drop=True),
    ], axis=1)

    scores_internal = pd.concat([
        data[["anonymous_id", "compositional_family", args.efa_col, "EFA_percentile", "EFA_quartile"]].reset_index(drop=True),
        score_df.reset_index(drop=True),
        X_raw[valid_features].reset_index(drop=True),
    ], axis=1)

    explained_df = pd.DataFrame({
        "component": pc_cols,
        "explained_variance": pca.explained_variance_,
        "explained_variance_ratio": pca.explained_variance_ratio_,
        "cumulative_explained_variance_ratio": np.cumsum(pca.explained_variance_ratio_),
    })

    loading_records = []
    for idx, c in enumerate(valid_features):
        rec = {
            "feature_column": c,
            "feature_label": PRETTY_LABELS.get(c, c),
            "feature_group": feature_group.get(c, "unknown"),
            "excel_block": "H-L" if feature_group.get(c) == "DFT" else "S-AK",
        }
        for j, pc in enumerate(pc_cols):
            val = float(components[j, idx])
            rec[f"{pc}_loading"] = val
            rec[f"abs_{pc}_loading"] = abs(val)
        loading_records.append(rec)
    loading_df = pd.DataFrame(loading_records)

    loading_long = []
    for _, row in loading_df.iterrows():
        for pc in pc_cols:
            loading_long.append({
                "component": pc,
                "feature_column": row["feature_column"],
                "feature_label": row["feature_label"],
                "feature_group": row["feature_group"],
                "excel_block": row["excel_block"],
                "loading": row[f"{pc}_loading"],
                "abs_loading": abs(row[f"{pc}_loading"]),
            })
    loading_long_df = pd.DataFrame(loading_long)

    pc_corr_records = []
    for pc in pc_cols:
        pc_corr_records.append(corr_pair(scores_public[pc], scores_public["EFA_percentile"], pc, "EFA_percentile"))
        pc_corr_records.append(corr_pair(scores_public[pc], data[args.efa_col], pc, "EFA_raw"))
    pc_corr_df = pd.concat(pc_corr_records, ignore_index=True).rename(columns={"x": "PC", "y": "target"})

    quartile_summary = (
        scores_public.groupby("EFA_quartile", dropna=False)[["PC1", "PC2", "EFA_percentile"]]
        .agg(["count", "mean", "std", "median", "min", "max"])
    )
    quartile_summary.columns = ["_".join([str(x) for x in col if x]) for col in quartile_summary.columns]
    quartile_summary = quartile_summary.reset_index()

    feature_info = pd.DataFrame({
        "feature_column": valid_features,
        "feature_label": [PRETTY_LABELS.get(c, c) for c in valid_features],
        "feature_group": [feature_group.get(c, "unknown") for c in valid_features],
        "excel_block": ["H-L" if feature_group.get(c) == "DFT" else "S-AK" for c in valid_features],
        "missing_count_before_imputation": [int(X_raw[c].isna().sum()) for c in valid_features],
        "median_used_for_imputation": [float(imputer.statistics_[i]) for i in range(len(valid_features))],
        "mean_after_imputation_before_scaling": [float(np.mean(X_imputed[:, i])) for i in range(len(valid_features))],
        "std_after_imputation_before_scaling": [float(np.std(X_imputed[:, i], ddof=0)) for i in range(len(valid_features))],
    })

    feature_group_summary = feature_info.groupby(["feature_group", "excel_block"]).agg(feature_count=("feature_column", "count")).reset_index()
    dropped_df = pd.DataFrame({"dropped_feature_column": dropped_features})
    orientation_df = pd.DataFrame(orientation_records)

    config_df = pd.DataFrame([
        {"parameter": "input_workbook", "value": str(input_path)},
        {"parameter": "sheet", "value": args.sheet},
        {"parameter": "DFT_feature_range", "value": f"{args.dft_start}-{args.dft_end}"},
        {"parameter": "chemistry_feature_range", "value": f"{args.chem_start}-{args.chem_end}"},
        {"parameter": "DFT_features", "value": ", ".join(dft_features)},
        {"parameter": "chemistry_features", "value": ", ".join(chem_features)},
        {"parameter": "valid_feature_count", "value": len(valid_features)},
        {"parameter": "rows_raw", "value": n_raw},
        {"parameter": "rows_used", "value": n_used},
        {"parameter": "PCA_input", "value": "H-L DFT descriptors + S-AK non-DFT chemistry descriptors"},
        {"parameter": "EFA_used_as_input", "value": "No; only used for coloring/evaluation"},
        {"parameter": "M_R_config_entropy_used", "value": "No"},
        {"parameter": "preprocessing", "value": "median imputation + standard scaling + PCA"},
        {"parameter": "EFA_percentile_definition", "value": "(rank - 1)/(N - 1), average ranks for ties"},
        {"parameter": "PC2_orientation", "value": "oriented positive with EFA percentile; PCA sign is arbitrary"},
    ])

    out_xlsx = outdir / "pca_efa_chaos_descriptor_space_H_L_S_AK_analysis.xlsx"
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        scores_public.to_excel(writer, sheet_name="PCA_Scores_Public", index=False)
        scores_internal.to_excel(writer, sheet_name="PCA_Scores_Internal", index=False)
        explained_df.to_excel(writer, sheet_name="PCA_Explained_Variance", index=False)
        loading_df.to_excel(writer, sheet_name="PCA_Loadings_Wide", index=False)
        loading_long_df.to_excel(writer, sheet_name="PCA_Loadings_Long", index=False)
        pc_corr_df.to_excel(writer, sheet_name="PC_EFA_Correlations", index=False)
        quartile_summary.to_excel(writer, sheet_name="EFA_Quartile_PC_Summary", index=False)
        feature_info.to_excel(writer, sheet_name="Feature_Info", index=False)
        feature_group_summary.to_excel(writer, sheet_name="Feature_Group_Summary", index=False)
        dropped_df.to_excel(writer, sheet_name="Dropped_Features", index=False)
        orientation_df.to_excel(writer, sheet_name="PC_Orientation", index=False)
        config_df.to_excel(writer, sheet_name="Config", index=False)
    autosize_workbook(out_xlsx)

    out_txt = outdir / "pca_efa_chaos_descriptor_space_H_L_S_AK_report.txt"
    write_report(
        out_txt=out_txt,
        input_path=input_path,
        sheet=args.sheet,
        n_raw=n_raw,
        n_used=n_used,
        dft_features=dft_features,
        chem_features=chem_features,
        valid_features=valid_features,
        dropped_features=dropped_features,
        explained_df=explained_df,
        pc_corr_df=pc_corr_df,
        loading_df=loading_df,
        out_xlsx=out_xlsx,
    )

    manifest = {
        "script": Path(__file__).name,
        "input": str(input_path),
        "sheet": args.sheet,
        "outputs": {"excel": str(out_xlsx), "txt": str(out_txt)},
        "PCA_input": "H-L + S-AK",
        "EFA_used_as_input": False,
    }
    (outdir / "pca_efa_chaos_descriptor_space_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print("Done. Outputs written to:")
    print(f"  Excel: {out_xlsx}")
    print(f"  TXT:   {out_txt}")


if __name__ == "__main__":
    main()
