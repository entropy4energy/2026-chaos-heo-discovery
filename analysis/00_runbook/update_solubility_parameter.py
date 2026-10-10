#!/usr/bin/env python3
"""Replace the solubility parameter of the locked cohort and rerun every
analysis that depends on the cohort (2026-10-05).

The solubility parameter of the 2026-09-25 package came from the archived
CHAOS export, where R* was computed with an earlier definition. The values
in ../solubility_parameter/solubility_parameter_503.csv use R* as the edge of
the cube with the volume per atom of each relaxed POCC supercell (see the
README in that folder). This script

1. copies that table into 00_cohort_lock/input/ and writes its
   solubility_parameter_1200K column into the solubility_parameter column of
   locked_493.xlsx (only those cells change), locked_493.csv,
   locked_493_typed.json (and the table record of
   locked_493.xlsx.inspect.ndjson while that file existed); the archived
   column stays in the copied table;
2. writes ../solubility_parameter/data_1002_delete_cf88_227_solubility_updated.xlsx,
   the prepared 503-row table of Figures 2 to 5 (raw_input/) with the same
   replacement, and the Figure 2b and 4b correlations of the new values;
3. reruns the archived scripts of folders 01 to 09 and 11 in place, on a copy
   of the new locked_493.xlsx in each input/, under legacy_gkf.py (same
   folder) so that the cross-validation folds are those of the production
   run; outputs, run_stdout.txt, run_stderr.txt and run_metadata.json are
   replaced, and absolute paths are written relative to the package, as in
   the 2026-09-25 deposit;
4. reassembles 12_rankings and 13_summary with assemble_production.py's
   functions (the manuscript and 503-row comparator inputs already in
   13_summary/input are kept; final_metrics.csv gains the Figure 6d Spearman
   coefficients of EFA with sigma_epsilon, d_bar, sigma_d and the solubility
   parameter, which the text quotes), writes the summary workbook with openpyxl in
   the layout of build_workbooks.mjs, and rewrites VALIDATION.json and
   MANIFEST.csv with the checks of verify_and_manifest.py.

Folder 10 reads only the Figure 6a scores and literature tables, which do
not use the solubility parameter, and is not rerun. The summary workbook's
preview image and inspection dump, which came from the tool that authored
the workbook, are removed rather than left stale.

    python3 update_solubility_parameter.py [COLUMN]

COLUMN defaults to solubility_parameter_1200K. With
solubility_parameter_archived the script reproduces the 2026-09-25 numbers,
which is how it was checked.
"""
from __future__ import annotations

import csv
import hashlib
import importlib.metadata
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parent
RUNBOOK = ROOT / "00_runbook"
TABLE = REPO / "solubility_parameter/solubility_parameter_503.csv"
LOCKED = ROOT / "00_cohort_lock/output/locked_493.xlsx"
ANALYSES = ["01_fig6a_compatibility", "02_fig6b_pca", "03_fig6c_si_quartiles", "04_fig6d_distortion",
            "05_fig6d_chemistry", "06_fig7a_regression", "07_fig7b_classification", "08_fig7c_enrichment",
            "09_fig7d_prioritization"]
TEXT = {".txt", ".json", ".csv", ".md"}
PACKAGES = {name: importlib.metadata.version(name)
            for name in ("numpy", "pandas", "scipy", "scikit-learn", "matplotlib", "openpyxl")}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for piece in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(piece)
    return h.hexdigest()


def relative(text: str) -> str:
    """Absolute paths inside the package, written as the 2026-09-25 deposit
    writes them (the package root as analysis)."""
    return text.replace(str(ROOT), "analysis")


def relative_files(folder: Path) -> None:
    for path in sorted(folder.rglob("*")):
        if not path.is_file():
            continue
        if path.suffix in TEXT:
            text = path.read_text(encoding="utf-8")
            if str(ROOT) in text:
                path.write_text(relative(text), encoding="utf-8")
        elif path.suffix == ".xlsx":
            wb = openpyxl.load_workbook(path)
            changed = False
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and str(ROOT) in cell.value:
                            cell.value = relative(cell.value)
                            changed = True
            if changed:
                wb.save(path)


def replace_column(path: Path, values: dict, sheet: str | None = None) -> int:
    """Only the solubility_parameter cells of the sheets with a name column change."""
    wb = openpyxl.load_workbook(path)
    n = 0
    for ws in wb.worksheets:
        if sheet and ws.title != sheet:
            continue
        head = [c.value for c in ws[1]]
        if "name" not in head or "solubility_parameter" not in head:
            continue
        cn, cs = head.index("name") + 1, head.index("solubility_parameter") + 1
        for r in range(2, ws.max_row + 1):
            name = ws.cell(row=r, column=cn).value
            if name is None:
                continue
            ws.cell(row=r, column=cs).value = values[name]
            n += 1
    wb.save(path)
    return n


def step_cohort(column: str) -> dict:
    table = pd.read_csv(TABLE)
    assert len(table) == 503 and table.name.is_unique and table[column].notna().all()
    shutil.copy2(TABLE, ROOT / "00_cohort_lock/input" / TABLE.name)
    values = dict(zip(table.name, table[column].astype(float)))
    out = ROOT / "00_cohort_lock/output"
    assert replace_column(LOCKED, values, "combined") == 493
    locked = pd.read_excel(LOCKED, sheet_name="combined")
    assert len(locked) == 493 and locked.columns.size == 44
    assert np.array_equal(locked.solubility_parameter.to_numpy(), locked.name.map(values).to_numpy())
    # the csv and typed json carry every other value as before
    old = pd.read_csv(out / "locked_493.csv")
    assert old.name.tolist() == locked.name.tolist()
    old["solubility_parameter"] = locked.solubility_parameter
    old.to_csv(out / "locked_493.csv", index=False, float_format="%.15g")
    # written as prepare_cohort.py wrote it (the values have at most 15
    # decimals, so double_precision=15 keeps them exactly)
    typed = json.loads((out / "locked_493_typed.json").read_text(encoding="utf-8"))
    typed = pd.DataFrame(typed["data"], columns=typed["columns"])
    typed["solubility_parameter"] = typed["name"].map(values)
    (out / "locked_493_typed.json").write_text(
        typed.to_json(orient="split", index=False, double_precision=15), encoding="utf-8")
    check = json.loads((out / "locked_493_typed.json").read_text(encoding="utf-8"))
    kk = check["columns"].index("solubility_parameter")
    assert [row[kk] for row in check["data"]] == locked.solubility_parameter.tolist()
    dump = out / "locked_493.xlsx.inspect.ndjson"
    lines = dump.read_text(encoding="utf-8").splitlines() if dump.exists() else []
    for j, line in enumerate(lines):
        record = json.loads(line)
        if record.get("kind") == "table" and record.get("sheet") == "combined":
            head = record["values"][0]
            kk, ii = head.index("solubility_parameter"), head.index("name")
            for row in record["values"][1:]:
                row[kk] = values[row[ii]]
            lines[j] = json.dumps(record, separators=(",", ":"), ensure_ascii=False)
    if dump.exists():
        dump.write_text("\n".join(lines), encoding="utf-8")
    record = {
        "date": "2026-10-05",
        "replaced_column": "solubility_parameter",
        "from_table": f"00_cohort_lock/input/{TABLE.name}",
        "table_column": column,
        "table_sha256": sha256(TABLE),
        "locked_493_xlsx_sha256": sha256(LOCKED),
        "rows_replaced": 493,
        "other_cells": "unchanged",
    }
    (out / "solubility_parameter_update.json").write_text(json.dumps(record, indent=2), encoding="utf-8")
    return values


def step_figures_2_to_5(values: dict) -> None:
    folder = TABLE.parent
    dest = folder / "data_1002_delete_cf88_227_solubility_updated.xlsx"
    shutil.copy2(REPO / "raw_input/data_1002_delete_cf88_227.xlsx", dest)
    assert replace_column(dest, values, "DFT properties") == 503
    dft = pd.read_excel(dest, sheet_name="DFT properties")
    chem = pd.read_excel(dest, sheet_name="ele descriptors")
    desc = {"lat_distortion_misfit_mean": "eps_bar", "lat_distortion_misfit_stdev": "sigma_eps",
            "lat_distortion_relax_distance_mean": "d_bar", "lat_distortion_relax_distance_stdev": "sigma_d",
            "solubility_parameter": "delta_s"}
    rows = []
    for c, label in desc.items():
        if c == "solubility_parameter":
            continue
        rows.append(("Figure 2b", label, dft.solubility_parameter.corr(dft[c]),
                     dft.solubility_parameter.corr(dft[c], method="spearman")))
    joined = dft[["name", "solubility_parameter"]].merge(chem, on="name", how="inner")
    assert len(joined) == 503
    for c in chem.columns:
        if c in ("name", "species"):
            continue
        rows.append(("Figure 4b", c, joined.solubility_parameter.corr(joined[c]),
                     joined.solubility_parameter.corr(joined[c], method="spearman")))
    pd.DataFrame(rows, columns=["figure", "with", "pearson_r", "spearman_rho"]).to_csv(
        folder / "correlations_figures_2b_4b.csv", index=False, float_format="%.6f")


def run(cmd: list, cwd: Path, env: dict):
    return subprocess.run(cmd, cwd=cwd, env=env, text=True, capture_output=True)


def step_analyses() -> None:
    shim = RUNBOOK / "legacy_gkf.py"
    mpl = Path(tempfile.mkdtemp(prefix="mpl_"))
    env = dict(os.environ, MPLBACKEND="Agg", MPLCONFIGDIR=str(mpl))
    for folder in ANALYSES:
        task = ROOT / folder
        metadata = json.loads((task / "run_metadata.json").read_text(encoding="utf-8"))
        code = task / "code"
        main = ROOT / metadata["script_snapshots"][0]["snapshot"]
        for item in metadata["script_snapshots"]:
            assert sha256(ROOT / item["snapshot"]) == item["sha256"], item["snapshot"]
        shutil.copy2(LOCKED, task / "input/locked_493.xlsx")
        output = task / "output"
        shutil.rmtree(output)
        output.mkdir()
        cmd = [sys.executable, str(shim), str(main), "--input", str(task / "input/locked_493.xlsx"),
               "--outdir", str(output)]
        print("RUN", folder, flush=True)
        done = run(cmd, code, env)
        (task / "run_stdout.txt").write_text(done.stdout, encoding="utf-8")
        (task / "run_stderr.txt").write_text(done.stderr, encoding="utf-8")
        metadata.update({
            "command": ["python3"] + cmd[1:],
            "returncode": done.returncode,
            "input_sha256": sha256(task / "input/locked_493.xlsx"),
            "python": sys.version,
            "packages": PACKAGES,
            "working_directory": str(code),
            "rerun": "2026-10-05, solubility parameter replaced (00_runbook/update_solubility_parameter.py)",
        })
        (task / "run_metadata.json").write_text(relative(json.dumps(metadata, indent=2)), encoding="utf-8")
        if done.returncode:
            print(done.stderr[-5000:], file=sys.stderr)
            raise RuntimeError(f"Failed: {folder}")
        relative_files(task)
    # folder 11: the script takes no arguments and writes into a task folder
    # that must not exist; it runs on a temporary task folder whose output is
    # copied into 11_model_robustness/output
    task = ROOT / "11_model_robustness"
    script = task / "code/run_model_robustness.py"
    tmp = Path(tempfile.mkdtemp(prefix="robust_"))
    runner = tmp / "runner.py"
    runner.write_text(
        "import importlib.util, pathlib\n"
        f"spec = importlib.util.spec_from_file_location('robust', {str(script)!r})\n"
        "m = importlib.util.module_from_spec(spec)\n"
        "spec.loader.exec_module(m)\n"
        f"m.TASK = pathlib.Path({str(tmp / 'task')!r})\n"
        f"m.LOCKED = pathlib.Path({str(LOCKED)!r})\n"
        f"m.COMMON_SOURCE = pathlib.Path({str(ROOT / '06_fig7a_regression/code/efa_screening_common.py')!r})\n"
        "m.main()\n", encoding="utf-8")
    print("RUN 11_model_robustness", flush=True)
    done = run([sys.executable, str(shim), str(runner)], task / "code", env)
    if done.returncode:
        print(done.stderr[-5000:], file=sys.stderr)
        raise RuntimeError("Failed: 11_model_robustness")
    assert sha256(tmp / "task/code/run_model_robustness.py") == sha256(script)
    assert sha256(tmp / "task/code/efa_screening_common.py") == sha256(task / "code/efa_screening_common.py")
    shutil.copy2(LOCKED, task / "input/locked_493.xlsx")
    shutil.rmtree(task / "output")
    shutil.copytree(tmp / "task/output", task / "output")
    (task / "run_stdout.txt").write_text(done.stdout, encoding="utf-8")
    (task / "run_stderr.txt").write_text(done.stderr, encoding="utf-8")
    (task / "run_metadata.json").write_text(json.dumps({
        "command": ["python3", "analysis/00_runbook/legacy_gkf.py", "RUNNER"],
        "runner": "imports code/run_model_robustness.py, sets TASK to a new temporary folder, LOCKED to "
                  "analysis/00_cohort_lock/output/locked_493.xlsx and COMMON_SOURCE to "
                  "analysis/06_fig7a_regression/code/efa_screening_common.py, and calls main(); the "
                  "temporary folder's output/ is this folder's output/",
        "returncode": done.returncode,
        "input_sha256": sha256(task / "input/locked_493.xlsx"),
        "python": sys.version,
        "packages": PACKAGES,
        "rerun": "2026-10-05, solubility parameter replaced (00_runbook/update_solubility_parameter.py)",
    }, indent=2), encoding="utf-8")
    relative_files(task)
    shutil.rmtree(tmp)
    shutil.rmtree(mpl)


def summary_workbook(path: Path, sheets: list) -> None:
    """build_workbooks.mjs summary layout, with openpyxl."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="263D56")
    for item in sheets:
        ws = wb.create_sheet(item["name"])
        ws.append(item["columns"])
        for row in item["data"]:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial")
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = head_fill
        ws.row_dimensions[1].height = 28
        for col, label in enumerate(item["columns"], start=1):
            label = str(label)
            width = (78 if label == "metric" else 58 if label == "analysis_folder" else 100 if label == "note"
                     else 70 if label == "suggested_display" else min(max(len(label) + 6, 20), 55))
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    wb.save(path)


DISPLAY = {
    "PC1 variance (%)": "{:.1f}%", "PC2 variance (%)": "{:.1f}%",
    "DFT R2": "{:.2f}", "DFT Spearman rho": "{:.2f}", "Chemistry R2": "{:.2f}",
    "Chemistry Spearman rho": "{:.2f}", "Combined R2": "{:.2f}", "Combined Spearman rho": "{:.2f}",
    "Chemistry ROC AUC": "{:.2f}", "Combined ROC AUC": "{:.2f}", "Chemistry top 10% hits": "{:.0f}/50",
    "Chemistry top 10% precision (%)": "{:.1f}%", "Chemistry top 10% enrichment": "{:.2f}-fold",
    "Predicted priority region size": "{:.0f}", "True-both precision (%)": "{:.1f}%",
    "True-both enrichment": "{:.2f}-fold",
}


def step_assemble() -> None:
    spec = importlib.util.spec_from_file_location("assemble_production", RUNBOOK / "assemble_production.py")
    ap = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ap)
    assert ap.ROOT == ROOT
    for folder in ("12_rankings", "13_summary"):
        shutil.copy2(LOCKED, ROOT / folder / "input/locked_493.xlsx")
    locked = pd.read_excel(LOCKED, sheet_name="combined")
    figure_index = ap.source_exports(locked)
    ranks = ap.rankings(locked)
    metrics = ap.metric_rows(locked)
    # the Results text also quotes the Figure 6d Spearman coefficients of
    # EFA with the other four distortion descriptors
    corr = ap.table("04_fig6d_distortion", "Correlation_Long")
    extra = []
    for descriptor, label in (("lat_distortion_misfit_stdev", "sigma_epsilon"),
                              ("lat_distortion_relax_distance_mean", "d_bar"),
                              ("lat_distortion_relax_distance_stdev", "sigma_d"),
                              ("solubility_parameter", "delta_s")):
        item = corr[corr.target_column.eq("EFA") & corr.descriptor_column.eq(descriptor) & corr.method.eq("Spearman")]
        assert len(item) == 1, descriptor
        extra.append(("Figure6d", f"Spearman_EFA_{label}", item.iloc[0].coefficient, "", "04_fig6d_distortion"))
    at = metrics.index[metrics.metric.eq("Spearman_EFA_epsilon_bar")][0] + 1
    metrics = pd.concat([metrics.iloc[:at], pd.DataFrame(extra, columns=metrics.columns), metrics.iloc[at:]],
                        ignore_index=True)
    metrics.to_csv(ROOT / "13_summary/output/final_metrics.csv", index=False)
    changes = ap.before_after(metrics)
    # the displayed values follow the production values; the ones the
    # replacement does not touch must come out as before
    archived = changes.suggested_display.copy()
    for i, q in enumerate(changes.quantity):
        if q in DISPLAY:
            changes.loc[i, "suggested_display"] = DISPLAY[q].format(changes.production_493[i])
    changes["note"] = changes["note"].where(
        changes.suggested_display.eq(archived),
        changes["note"] + " Production value and display updated 2026-10-05 for the solubility parameter.")
    changes.to_csv(ROOT / "13_summary/output/before_after_manuscript_numbers.csv", index=False)
    literature = pd.read_csv(ROOT / "10_literature_candidates/output/literature_table_43_retained.csv")
    sensitivity = pd.read_csv(ROOT / "11_model_robustness/output/twenty_seed_metrics.csv")
    sheets = [ap.workbook_sheet("Final_Metrics", metrics), ap.workbook_sheet("Before_After", changes),
              ap.workbook_sheet("Top_50_Combined", ranks.head(50).drop(columns=["EFA_observed_internal"])),
              ap.workbook_sheet("Twenty_Seeds", sensitivity),
              ap.workbook_sheet("Figure_Sources", figure_index),
              ap.workbook_sheet("Literature_43", literature.drop(columns=["source_line_text"]))]
    (ROOT / "13_summary/output/summary_workbook.json").write_text(
        json.dumps({"sheets": sheets}, indent=2, ensure_ascii=False, allow_nan=False), encoding="utf-8")
    summary_workbook(ROOT / "13_summary/output/CHAOS_EFA493_production_summary.xlsx", sheets)
    for stale in ("summary_preview.png", "CHAOS_EFA493_production_summary.xlsx.inspect.ndjson"):
        (ROOT / "13_summary/output" / stale).unlink(missing_ok=True)
    for folder in ("12_rankings", "13_summary"):
        relative_files(ROOT / folder)


def step_verify() -> None:
    """verify_and_manifest.py's checks; the archived script sources are
    outside the package (placeholder paths), so each snapshot is checked
    against its recorded hash only."""
    cohort = pd.read_excel(LOCKED, sheet_name="combined")
    assert len(cohort) == 493 and cohort.full_cation_key.is_unique
    assert cohort.cation_count.eq(4).all() and cohort.match_count.eq(1).all()
    assert cohort.Metal5.isna().all()
    assert len(pd.read_csv(ROOT / "00_cohort_lock/output/excluded_ten_quinaries.csv")) == 10
    input_hash = sha256(LOCKED)
    checks = []
    for folder in sorted(ROOT.glob("0[1-9]_*")) + [ROOT / "11_model_robustness"]:
        metadata = json.loads((folder / "run_metadata.json").read_text())
        assert metadata["returncode"] == 0 and metadata["input_sha256"] == input_hash, folder
        assert sha256(folder / "input/locked_493.xlsx") == input_hash
        if folder.name != "11_model_robustness":
            assert (folder / "output/figure_source_csv").is_dir()
            for item in metadata["script_snapshots"]:
                assert sha256(ROOT / item["snapshot"]) == item["sha256"]
        checks.append({"folder": folder.name, "status": "passed"})
    for folder in ("12_rankings", "13_summary"):
        assert sha256(ROOT / folder / "input/locked_493.xlsx") == input_hash
    ranks = pd.read_csv(ROOT / "12_rankings/output/all_493_ranked.csv")
    metrics = pd.read_csv(ROOT / "13_summary/output/final_metrics.csv")
    changes = pd.read_csv(ROOT / "13_summary/output/before_after_manuscript_numbers.csv")
    assert len(ranks) == 493 and ranks.full_cation_key.is_unique
    assert ranks.high_EFA_observed.sum() == 124
    assert ranks.predicted_priority_region.sum() == 52
    combined = pd.read_csv(ROOT / "12_rankings/output/combined_top_10_percent_50.csv")
    chemistry = pd.read_csv(ROOT / "12_rankings/output/chemistry_top_10_percent_50.csv")
    assert len(combined) == 50 and len(chemistry) == 50 and int(chemistry.high_EFA_observed.sum()) == 42
    priority = pd.read_csv(ROOT / "12_rankings/output/figure7d_priority_region_52.csv")
    joint = "true high-EFA + high descriptor compatibility score"
    assert len(priority) == 52 and int(priority.observed_priority_class.eq(joint).sum()) == 48
    assert len(pd.read_csv(ROOT / "10_literature_candidates/output/literature_table_43_retained.csv")) == 43
    assert len(pd.read_csv(ROOT / "10_literature_candidates/output/literature_table_4_removed.csv")) == 4
    assert len(pd.read_csv(ROOT / "11_model_robustness/output/twenty_seed_metrics.csv")) == 20
    assert json.loads((ROOT / "11_model_robustness/output/feature_group_check.json").read_text())["numerically_identical_at_1e_10"]
    assert len(metrics) >= 60 and len(changes) == 18
    assert len(pd.read_csv(ROOT / "13_summary/output/figure_source_index.csv")) == 25
    summary = pd.ExcelFile(ROOT / "13_summary/output/CHAOS_EFA493_production_summary.xlsx")
    assert {"Final_Metrics", "Before_After", "Top_50_Combined", "Twenty_Seeds", "Figure_Sources",
            "Literature_43"}.issubset(summary.sheet_names)
    for path in ROOT.rglob("*"):
        if path.is_file() and path.suffix in TEXT:
            assert str(ROOT) not in path.read_text(encoding="utf-8"), path
    result = {"cohort_rows": len(cohort), "excluded_rows": 10, "high_EFA": 124,
              "analysis_checks": checks, "chemistry_top_50_hits": 42,
              "combined_top_50_hits": int(combined.high_EFA_observed.sum()),
              "priority_region": {"total": 52, "true_joint_class": 48},
              "literature_rows": {"retained": 43, "removed": 4},
              "seed_count": 20, "figure_source_csv_count": 25,
              "solubility_parameter_update": "2026-10-05, 00_cohort_lock/output/solubility_parameter_update.json"}
    (ROOT / "VALIDATION.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    manifest = ROOT / "MANIFEST.csv"
    rows = []
    for path in sorted(ROOT.rglob("*")):
        if not path.is_file() or path.is_symlink() or path == manifest or "__pycache__" in path.parts:
            continue
        rel = path.relative_to(ROOT)
        role = ("runbook/metadata" if rel.parts[0] == "00_runbook" or len(rel.parts) == 1 else
                "analysis input" if "input" in rel.parts else
                "script snapshot" if "code" in rel.parts else "analysis output")
        rows.append((str(rel), role, path.stat().st_size, sha256(path)))
    rows.append(("MANIFEST.csv", "self-reference; hash intentionally omitted", "", ""))
    with manifest.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(("path_relative_to_package", "role", "bytes", "sha256"))
        writer.writerows(rows)
    print("validated", len(cohort), "rows; manifest files", len(rows))


def main() -> None:
    column = sys.argv[1] if len(sys.argv) > 1 else "solubility_parameter_1200K"
    values = step_cohort(column)
    step_figures_2_to_5(values)
    step_analyses()
    step_assemble()
    for path in ROOT.rglob("__pycache__"):
        shutil.rmtree(path)
    step_verify()


if __name__ == "__main__":
    main()
