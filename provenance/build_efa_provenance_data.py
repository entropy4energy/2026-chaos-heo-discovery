#!/usr/bin/env python3
"""Read-only reconstruction checks and records for EFA provenance workbook.

This is a new audit/packaging script. It is not the historical MATLAB merge.
"""

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


FILES = [
    ("data_1002_delete_cf88_227.xlsx", "Prepared rocksalt DFT and element-descriptor export"),
    ("extract_element.m", "Original parser that keeps the first four cations"),
    ("extracted_metals.xlsx", "Four extracted cation columns"),
    ("chaos_data.xlsx", "Row-aligned DFT, chemistry, and four-metal input to merge"),
    ("lib5_organized_072125.json", "Four-cation EFA records; filename date is unverified"),
    ("lib5_efa_deed.xlsx", "Flattened EFA raw and cleaned sheets"),
    ("merge_efa_deed_data.m", "Original four-metal merge script"),
    ("chaos_data_with_efa_deed.xlsx", "Historical merged output"),
]
EFA_FIELDS = ("EFA",)


def sheet_rows(path, sheet):
    return list(load_workbook(path, read_only=True, data_only=True)[sheet].values)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def cations(name):
    composition = name.split(":", 1)[0]
    tokens = re.findall(r"[A-Z][a-z]?(?:_pv|_sv)?", composition)
    return [re.match(r"[A-Z][a-z]?", t).group() for t in tokens if not t.startswith("O")]


def equal(x, y):
    if x == y or x in (None, "") and y in (None, ""):
        return True
    return isinstance(x, (float, int)) and isinstance(y, (float, int)) and abs(x - y) < 1e-10


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source-dir", type=Path, required=True)
    ap.add_argument("--output", type=Path, required=True)
    args = ap.parse_args()
    root = args.source_dir
    prepared = load_workbook(root / FILES[0][0], read_only=True, data_only=True)
    dft = list(prepared["DFT_properties"].values)
    chemistry = list(prepared["ele_descriptors"].values)
    metals = sheet_rows(root / "extracted_metals.xlsx", "Sheet1")
    joined_input = sheet_rows(root / "chaos_data.xlsx", "Sheet1")
    efa_wb = load_workbook(root / "lib5_efa_deed.xlsx", read_only=True, data_only=True)
    efa_raw = list(efa_wb["raw"].values)
    efa_cleaned = list(efa_wb["cleaned"].values)
    combined = sheet_rows(root / "chaos_data_with_efa_deed.xlsx", "combined")
    organized = json.loads((root / "lib5_organized_072125.json").read_text())

    assert all(len(t) == 504 for t in (dft, chemistry, metals, joined_input, combined))
    assert len(efa_raw) == len(efa_cleaned) == 530 and len(organized) == 529
    assert len(set(tuple(sorted(row[:4])) for row in efa_cleaned[1:])) == 529
    assert all(len(x["Elements"]) == 4 for x in organized)

    dft_header = {field: i for i, field in enumerate(dft[0])}
    chem_header = {field: i for i, field in enumerate(chemistry[0])}
    input_header = list(joined_input[0])
    output_header = list(combined[0])
    parent_by_key = {}
    quaternary = []
    quinary = []
    checks = Counter()
    for row_number, (dft_row, chem_row, metal_row, input_row, output_row) in enumerate(
        zip(dft[1:], chemistry[1:], metals[1:], joined_input[1:], combined[1:]), 2
    ):
        name = output_row[0]
        assert name == dft_row[0] == chem_row[0] == input_row[0]
        checks["name_rows_matched"] += 1
        all_cations = cations(name)
        assert tuple(all_cations[:4]) == tuple(metal_row[:4]) == tuple(output_row[1:5])
        checks["extracted_metals_matched"] += 1
        assert all(equal(a, b) for a, b in zip(input_row[:37], output_row[:37]))
        checks["merged_input_first_37_matched"] += 1
        for field in input_header[5:37]:
            i = input_header.index(field)
            source = (dft_row[dft_header[field]] if field in dft_header
                      else chem_row[chem_header[field]])
            assert equal(source, input_row[i]), (row_number, field)
            checks["source_descriptor_cells_matched"] += 1
        key = tuple(sorted(output_row[1:5]))
        source_row = int(output_row[42])
        clean = efa_cleaned[source_row - 1]
        raw = efa_raw[source_row - 1]
        record = organized[source_row - 2]
        assert key == tuple(sorted(clean[:4]))
        assert output_row[43] == 1
        assert all(equal(output_row[37 + j], clean[4 + j]) for j in range(len(EFA_FIELDS)))
        assert tuple(record["Elements"]) == tuple(clean[:4])
        assert tuple(record["species"]) == tuple(raw[:5])
        assert all(abs(record[field] - clean[4 + j]) < 1e-10
                   for j, field in enumerate(EFA_FIELDS))
        checks["efa_matches_source"] += 1

        if len(all_cations) == 4 and "0.25x" in name:
            item = {
                "combined_row": row_number,
                "pocc_parent_name": name,
                "Metal1": output_row[1], "Metal2": output_row[2],
                "Metal3": output_row[3], "Metal4": output_row[4],
                "sorted_four_metal_key": ",".join(key),
                "dft_source_row": row_number,
                "element_descriptor_row": row_number,
                "chaos_data_row": row_number,
                "efa_cleaned_row": source_row,
                "json_record_index_1_based": source_row - 1,
                "chem_id": record["chem_id"],
                "AUID": None,
                "AUID_status": "not present in local sources",
                "source_configuration_count": record["dg_len"],
                "configuration_review": "review count below modal 18" if record["dg_len"] < 18 else "",
                "ensemble_publication_status": "not recorded",
                "database_snapshot_or_export_date": "not recorded",
            }
            quaternary.append(item)
            assert key not in parent_by_key
            parent_by_key[key] = (row_number, item)
        elif len(all_cations) == 5 and "0.2x" in name:
            quinary.append({
                "combined_row": row_number,
                "pocc_name": name,
                "five_cations": ",".join(all_cations),
                "fifth_cation_lost_by_parser": all_cations[4],
                "four_metal_join_key": ",".join(key),
                "efa_cleaned_row": source_row,
                "assigned_EFA": output_row[37],
                "correct_five_cation_EFA": None,
            })
        else:
            raise AssertionError((row_number, name, all_cations))

    assert len(quaternary) == 493 and len(quinary) == 10
    assert len(set(x["efa_cleaned_row"] for x in quaternary)) == 493
    assert sum(x["source_configuration_count"] < 18 for x in quaternary) == 38
    for item in quinary:
        parent_row, parent = parent_by_key[tuple(item["four_metal_join_key"].split(","))]
        assert item["efa_cleaned_row"] == parent["efa_cleaned_row"]
        item["four_cation_parent_combined_row"] = parent_row
        item["four_cation_parent_name"] = parent["pocc_parent_name"]

    sources = [
        {"file": name, "role": role, "sha256": sha256(root / name),
         "database_snapshot_or_export_date": "not recorded in file"}
        for name, role in FILES
    ]
    payload = {
        "summary": {
            "prepared_rocksalt_rows": 503,
            "quaternary_rows": 493,
            "quinary_rows_misjoined": 10,
            "efa_source_rows": 529,
            "unused_efa_source_rows": 36,
            "quaternary_rows_with_configuration_count_below_18": 38,
            "records_with_AUID_in_this_folder": 0,
            "verified_database_snapshot_or_export_date": None,
        },
        "checks": dict(checks),
        "quaternary": quaternary,
        "quinary": quinary,
        "sources": sources,
    }
    args.output.write_text(json.dumps(payload, indent=2))
    print(json.dumps({"summary": payload["summary"], "checks": payload["checks"]}, indent=2))


if __name__ == "__main__":
    main()
